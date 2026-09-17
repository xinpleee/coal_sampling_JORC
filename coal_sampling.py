"""Coal geology automatic sampling engine.

The engine is deliberately independent from tkinter so it can be tested and used
from scripts. It never writes to the input workbook.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

EPS = Decimal("0.000001")


@dataclass
class Parameters:
    thin_max: Decimal = Decimal("0.05")
    ordinary_max: Decimal = Decimal("0.20")
    group_gap_max: Decimal = Decimal("2.00")
    board_section_co_max: Decimal = Decimal("3.00")
    coal_layer_no_sample_max: Decimal = Decimal("0.20")
    board_length: Decimal = Decimal("0.20")
    coal_sample_max: Decimal = Decimal("1.00")
    sample_thin_max: Decimal = Decimal("0.05")
    thickness_tolerance: Decimal = Decimal("0.001")
    eps: Decimal = EPS
    special_codes: tuple[str, ...] = ("XM", "ZM", "CO-PY", "XT")
    ordinary_codes: tuple[str, ...] = ("CG", "GV", "SS", "ST", "MS", "BMS")

    def jsonable(self):
        result = asdict(self)
        for key, value in result.items():
            if isinstance(value, Decimal):
                result[key] = str(value)
            elif isinstance(value, tuple):
                result[key] = list(value)
        return result


@dataclass
class Record:
    source_sheet: str
    source_row: int
    from_: Decimal | None
    to: Decimal | None
    raw_thickness: Decimal | None
    lithology: str
    calculated: Decimal | None = None
    classification: str = ""
    coal_layer: str = ""
    coal_section: str = ""
    coal_group: str = ""
    issue: str = ""
    sampled: bool = False

    @property
    def valid_depth(self):
        return self.from_ is not None and self.to is not None and self.to > self.from_

    @property
    def thickness(self):
        return self.calculated or Decimal(0)


@dataclass
class Sample:
    from_: Decimal
    to: Decimal
    lithology: str
    sample_type: str
    coal_group: str = ""
    coal_section: str = ""
    coal_layer: str = ""
    comments: str = ""
    sources: list[str] = field(default_factory=list)
    sample_no: str = ""

    @property
    def thickness(self):
        return self.to - self.from_


def dec(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    try:
        value = Decimal(str(value).strip())
        return value if value.is_finite() else None
    except (InvalidOperation, ValueError, TypeError):
        return None


def normalize_lithology(value: Any) -> str:
    text = "" if value is None else str(value).strip().upper()
    for dash in ("－", "—", "–", "−", "﹣"):
        text = text.replace(dash, "-")
    return "".join(text.split())


class SamplingEngine:
    def __init__(self, params: Parameters | None = None):
        self.params = params or Parameters()
        self.records: list[Record] = []
        self.samples: list[Sample] = []
        self.reviews: list[dict[str, Any]] = []
        self.selected_special: set[str] = set()
        self.selected_thin_coal: set[str] = set()

    def _resolve_sheet_name(self, sheet_names: list[str]) -> str:
        exact = "Litho【煤岩编录】"
        if exact in sheet_names:
            return exact
        candidates = [n for n in sheet_names if "litho" in n.lower()]
        if candidates:
            if len(candidates) > 1:
                raise ValueError("存在多个候选工作表，请保留唯一的 Litho 或 LOG 工作表")
            return candidates[0]
        candidates = [n for n in sheet_names if n.lower() == "log"]
        if candidates:
            if len(candidates) > 1:
                raise ValueError("存在多个候选工作表，请保留唯一的 Litho 或 LOG 工作表")
            return candidates[0]
        raise ValueError("未找到 Litho【煤岩编录】、Litho 或 LOG 工作表")

    def load_xlsx(self, filename: str | Path):
        """Read the selected sheet twice, preferring cached formula values."""
        import openpyxl

        path = Path(filename)
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
        values = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheet_name = self._resolve_sheet_name(formulas.sheetnames)
        fs, vs = formulas[sheet_name], values[sheet_name]
        is_litho = "litho" in sheet_name.lower() or sheet_name == "Litho【煤岩编录】"
        start, cols = (5, (1, 2, 3, 6)) if is_litho else (2, (6, 7, 8, 11))

        self.records = []
        self.reviews = []
        for row_no in range(start, fs.max_row + 1):
            try:
                frow = next(fs.iter_rows(min_row=row_no, max_row=row_no, values_only=False))
                vrow = next(vs.iter_rows(min_row=row_no, max_row=row_no, values_only=True))
            except StopIteration:
                continue

            vals = []
            for col in cols:
                cell = frow[col - 1]
                cached = vrow[col - 1]
                is_formula = isinstance(cell.value, str) and cell.value.startswith("=")
                if is_formula and cached is None:
                    vals.append(None)
                else:
                    vals.append(cached if is_formula else cell.value)

            fr, to, raw, lith = vals
            record = Record(sheet_name, row_no, dec(fr), dec(to), dec(raw), normalize_lithology(lith))
            if not record.valid_depth:
                record.issue = "INVALID_DEPTH_OR_EMPTY_FIELD"
                self.reviews.append(self.review(record, record.issue))
            else:
                record.calculated = record.to - record.from_
                if record.raw_thickness is None or abs(record.raw_thickness - record.calculated) > self.params.thickness_tolerance:
                    record.issue = "RAW_THICKNESS_MISMATCH"
                    self.reviews.append(self.review(record, record.issue))
                if not record.lithology:
                    record.issue = "EMPTY_LITHOLOGY"
                    self.reviews.append(self.review(record, record.issue))
            self.records.append(record)

        self.records.sort(key=lambda x: (x.from_ is None, x.from_ or Decimal(0), x.to or Decimal(0), x.source_row))
        self.calculate()

    def review(self, r, reason):
        return {
            "Source Sheet": r.source_sheet,
            "Source Row": r.source_row,
            "From": r.from_,
            "To": r.to,
            "Thickness": r.thickness,
            "Lithology": r.lithology,
            "Reason": reason,
        }

    def classify(self, r: Record):
        if not r.valid_depth or not r.lithology:
            return "INVALID"
        if r.lithology == "CO":
            return "CO"
        if r.thickness <= self.params.thin_max:
            return "THIN"
        if r.lithology in self.params.special_codes:
            return "SPECIAL"
        if r.lithology in self.params.ordinary_codes:
            return "ORDINARY_THIN" if r.thickness <= self.params.ordinary_max else "ORDINARY"
        return "SPECIAL"

    def calculate(self):
        self.samples = []
        self.reviews = [
            x for x in self.reviews if x.get("Reason") in ("INVALID_DEPTH_OR_EMPTY_FIELD", "RAW_THICKNESS_MISMATCH", "EMPTY_LITHOLOGY")
        ]

        for r in self.records:
            r.classification = self.classify(r)
            r.coal_layer = ""
            r.coal_section = ""
            r.coal_group = ""
            r.sampled = False

        valid = [r for r in self.records if r.valid_depth and r.lithology]
        self._check_geometry(valid)
        self._hierarchy(valid)
        self._thin_coal_specials(valid)
        self._independent_layers(valid)
        self._co_units(valid)
        self._boards(valid)
        self._normalize()
        self._make_reviews(valid)

    def _check_geometry(self, rs):
        covered = None
        for r in rs:
            if covered is not None:
                if r.from_ > covered + self.params.eps:
                    self.reviews.append(self.review(r, "DEPTH_GAP"))
                elif r.from_ < covered - self.params.eps:
                    self.reviews.append(self.review(r, "UNRESOLVED_OVERLAP"))
            covered = max(covered or r.to, r.to)

    def _hierarchy(self, rs):
        coal_layers = []
        current = []
        prev = None
        for r in rs:
            if r.classification == "CO":
                if prev is not None and abs(r.from_ - prev.to) > self.params.eps:
                    current = []
                current.append(r)
                prev = r
            else:
                if current:
                    coal_layers.append(current)
                    current = []
                prev = r
        if current:
            coal_layers.append(current)

        for layer_index, layer in enumerate(coal_layers, 1):
            layer_id = f"CL-{layer_index:03d}"
            for r in layer:
                r.coal_layer = layer_id

        usable = {"CO", "THIN", "SPECIAL", "ORDINARY_THIN"}
        sections = []
        block = []
        prev = None
        for r in rs:
            contiguous = prev is not None and abs(r.from_ - prev.to) <= self.params.eps
            if r.classification in usable and contiguous:
                block.append(r)
            else:
                if any(x.classification == "CO" for x in block):
                    sections.append(block)
                block = [r] if r.classification in usable else []
            prev = r
        if any(x.classification == "CO" for x in block):
            sections.append(block)

        for section_index, section in enumerate(sections, 1):
            section_id = f"CS-{section_index:03d}"
            for r in section:
                r.coal_section = section_id

        group_no = 0
        end = None
        for section in sections:
            start, stop = section[0].from_, section[-1].to
            if end is None or start - end > self.params.group_gap_max:
                group_no += 1
            for r in section:
                r.coal_group = f"CG-{group_no:03d}"
            end = stop

    def _thin_coal_specials(self, rs):
        for idx, r in enumerate(rs):
            if not r.coal_layer or r.thickness >= self.params.coal_layer_no_sample_max:
                continue

            neighbours = []
            if idx and rs[idx - 1].classification == "SPECIAL":
                neighbours.append(rs[idx - 1])
            if idx + 1 < len(rs) and rs[idx + 1].classification == "SPECIAL":
                neighbours.append(rs[idx + 1])
            if not neighbours:
                self.reviews.append(self.review(r, "THIN_COAL_NO_ADJACENT_SPECIAL"))
                continue

            if len(neighbours) == 2 and neighbours[0].lithology == neighbours[1].lithology:
                self.samples.append(
                    self._sample(neighbours[0].from_, neighbours[1].to, neighbours[0].lithology, "特殊夹层样", r, [neighbours[0], r, neighbours[1]])
                )
                r.sampled = True
                for item in neighbours:
                    item.sampled = True
                continue

            target = neighbours[-1]
            self.samples.append(
                self._sample(
                    r.from_ if target.from_ > r.from_ else target.from_,
                    target.to if target.from_ > r.from_ else r.to,
                    target.lithology,
                    "特殊夹层样",
                    r,
                    [r, target],
                )
            )
            r.sampled = True
            target.sampled = True

    def _independent_layers(self, rs):
        for r in rs:
            if r.sampled:
                continue
            if r.classification == "SPECIAL" and r.thickness > self.params.thin_max:
                self.samples.append(self._sample(r.from_, r.to, r.lithology, "特殊夹层样", r, [r]))
                r.sampled = True
            elif r.classification == "ORDINARY_THIN":
                self.samples.append(self._sample(r.from_, r.to, r.lithology, "普通夹层样", r, [r]))
                r.sampled = True

    def _co_units(self, rs):
        unit = []

        def flush():
            nonlocal unit
            if unit:
                self._split_co(unit)
                unit = []

        for idx, r in enumerate(rs):
            accepted = r.classification == "CO" and not r.sampled
            if r.classification == "THIN":
                accepted = idx > 0 and idx + 1 < len(rs) and rs[idx - 1].classification == "CO" and rs[idx + 1].classification == "CO"
            if accepted:
                unit.append(r)
            else:
                flush()
        flush()

    def _split_co(self, unit):
        total = unit[-1].to - unit[0].from_
        if total <= Decimal("1.5"):
            chunks = [unit]
        else:
            chunks = []
            start = 0
            accumulated = Decimal(0)
            thin = Decimal(0)
            for idx, r in enumerate(unit):
                if accumulated + r.thickness > self.params.coal_sample_max and idx > start:
                    chunks.append(unit[start:idx])
                    start = idx
                    accumulated = thin = Decimal(0)
                accumulated += r.thickness
                if r.classification == "THIN":
                    thin += r.thickness
                if thin > self.params.sample_thin_max and idx > start:
                    chunks.append(unit[start:idx])
                    start = idx
                    accumulated = r.thickness
                    thin = r.thickness if r.classification == "THIN" else Decimal(0)
            if start < len(unit):
                chunks.append(unit[start:])

        for chunk in chunks:
            if chunk[-1].to - chunk[0].from_ <= 0:
                continue
            self.samples.append(self._sample(chunk[0].from_, chunk[-1].to, "CO", "煤样", chunk[0], chunk))
            for r in chunk:
                r.sampled = True

    def _sample(self, fr, to, lith, typ, owner, parts):
        extras = [
            f"{int((r.thickness * 100).to_integral_value())}cm,{r.lithology}"
            for r in parts if r.classification != "CO" and r is not owner
        ]
        return Sample(
            fr,
            to,
            lith,
            typ,
            owner.coal_group,
            owner.coal_section,
            owner.coal_layer,
            "; ".join(extras),
            [f"{r.source_sheet}:{r.source_row}" for r in parts],
        )

    def _boards(self, rs):
        for group in sorted({r.coal_group for r in rs if r.coal_group}):
            members = [r for r in rs if r.coal_group == group]
            co_total = sum((r.thickness for r in members if r.classification == "CO"), Decimal(0))
            if co_total <= self.params.board_section_co_max:
                continue
            first, last = members[0], members[-1]
            self.samples.append(
                Sample(
                    max(Decimal(0), first.from_ - self.params.board_length),
                    first.from_,
                    "RF-" + first.lithology,
                    "顶板样",
                    group,
                    first.coal_section,
                    first.coal_layer,
                )
            )
            self.samples.append(
                Sample(
                    last.to,
                    last.to + self.params.board_length,
                    "FL-" + last.lithology,
                    "底板样",
                    group,
                    last.coal_section,
                    last.coal_layer,
                )
            )

    def _normalize(self):
        self.samples = [s for s in self.samples if s.to > s.from_]
        self.samples.sort(key=lambda s: (s.from_, s.to, s.sample_type))

        merged = []
        for s in self.samples:
            if (
                merged
                and s.lithology == merged[-1].lithology
                and s.sample_type == merged[-1].sample_type
                and abs(s.from_ - merged[-1].to) <= self.params.eps
                and s.sample_type != "煤样"
            ):
                merged[-1].to = s.to
                merged[-1].comments = "; ".join(x for x in (merged[-1].comments, s.comments) if x)
                merged[-1].sources.extend(s.sources)
            else:
                merged.append(s)

        self.samples = merged
        for index, s in enumerate(self.samples, 1):
            s.sample_no = f"S{index:04d}"

    def _make_reviews(self, rs):
        for r in rs:
            if r.classification == "SPECIAL" and not r.sampled:
                self.reviews.append(self.review(r, "SPECIAL_NOT_SAMPLED"))
            if r.coal_layer and r.thickness < self.params.coal_layer_no_sample_max and not r.sampled:
                self.reviews.append(self.review(r, "THIN_COAL_NOT_SAMPLED"))

    def result_rows(self):
        return [
            {
                "Sample No": s.sample_no,
                "煤组": s.coal_group,
                "煤段": s.coal_section,
                "煤层": s.coal_layer,
                "From": float(s.from_),
                "To": float(s.to),
                "Thickness": float(s.thickness),
                "Lithology": s.lithology,
                "Sample Type": s.sample_type,
                "Comments": s.comments,
                "Sources": "; ".join(s.sources),
            }
            for s in self.samples
        ]

    def review_rows(self):
        return self.reviews

    def original_rows(self):
        out = []
        for r in self.records:
            state = "不取样"
            hits = [s.sample_no for s in self.samples if s.from_ < (r.to or Decimal(0)) - self.params.eps and s.to > (r.from_ or Decimal(0)) + self.params.eps]
            if hits:
                state = "、".join(hits)
            out.append(
                {
                    "Source Row": r.source_row,
                    "From": r.from_,
                    "To": r.to,
                    "Thickness": r.thickness,
                    "Lithology": r.lithology,
                    "Classification": r.classification,
                    "煤层": r.coal_layer,
                    "煤段": r.coal_section,
                    "煤组": r.coal_group,
                    "采样状态": state,
                    "Issue": r.issue,
                }
            )
        return out
